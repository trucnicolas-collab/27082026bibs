"""
Test: le PPTX généré ne doit JAMAIS contenir de fichiers ZIP dupliqués.

Contexte régression (01/09/2026) : après suppression de slides semaines
inutilisées puis insertion du slide "Accès et logistique", python-pptx
attribuait au nouveau slide un partname déjà utilisé (slide18.xml collision).
Résultat : PPTX corrompu, PowerPoint refuse d'ouvrir ("PowerPoint a détecté
un problème dans le contenu").

Fix : _add_slide_safe_partname() alloue un partname unique en scannant
les partnames existants dans le package.
"""
import sys
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import pptx_export  # noqa: E402
from pptx import Presentation  # noqa: E402


def _aggregate_fn_stub(_doc):
    return {
        "nuit_es": {1: {"date": "01/01/2026", "sr": "PGC", "allees_str": "1",
                         "eeg": 100, "rails_es": 0, "sa": 0, "cam": 0}},
        "nuit_cam": {},
        "totals_by_nuit": {1: {"eeg": 100, "cam": 0, "sa": 0}},
        "dates_map": {"1": "01/01/2026"},
        "weeks": [1],
        "all_nights": [1],
        "cam_nights": [],
        "hide_sa_mag": False,
    }


def _minimal_dataset():
    return {
        "raw_records": [
            {"Secteur": "PGC", "Rayon": "Boissons", "N° allée": 1,
             "Type": "EEG", "Désignation": "ES 1.5 (noir)", "Quantité": 100},
        ],
        "filename": "test.xlsx", "store_name": "TEST", "store_code": "H0000",
        "phasage": {"es": {"nb_nuits": 3, "rows": []},
                    "cam": {"nb_nuits": 2, "rows": []}},
        "recap_rows": [],
    }


def _check_no_zip_duplicates(pptx_bytes):
    with zipfile.ZipFile(BytesIO(pptx_bytes)) as z:
        names = z.namelist()
    counts = Counter(names)
    dups = {k: v for k, v in counts.items() if v > 1}
    assert not dups, f"PPTX ZIP contient des entrées dupliquées : {dups}"


def test_pptx_no_duplicate_slide_partnames_short_week():
    """1 semaine → 3 slides semaine supprimées + logistique inséré = combo à risque."""
    d = _minimal_dataset()
    data = pptx_export.build_pptx(
        d, aggregate_fn=_aggregate_fn_stub, recap_rows=[],
        summary={"allees": [], "totals": {}}, detail_cam_rows=[],
    )
    _check_no_zip_duplicates(data)
    # Le PPTX doit être ouvrable
    p = Presentation(BytesIO(data))
    assert len(p.slides) >= 15


def test_pptx_no_duplicate_slide_partnames_two_weeks():
    """2 semaines → 2 slides semaine supprimées."""
    def agg(_):
        return {
            "nuit_es": {1: {"date": "", "sr": "", "allees_str": "1",
                            "eeg": 100, "rails_es": 0, "sa": 0, "cam": 0}},
            "nuit_cam": {},
            "totals_by_nuit": {1: {"eeg": 100, "cam": 0, "sa": 0}},
            "dates_map": {"1": ""},
            "weeks": [2, 3],
            "all_nights": [1, 2, 3, 4, 5],
            "cam_nights": [],
            "hide_sa_mag": False,
        }
    d = _minimal_dataset()
    data = pptx_export.build_pptx(
        d, aggregate_fn=agg, recap_rows=[],
        summary={"allees": [], "totals": {}}, detail_cam_rows=[],
    )
    _check_no_zip_duplicates(data)


def test_pptx_no_duplicate_slide_partnames_five_weeks():
    """5 semaines → clone d'un slide semaine + éventuelles suppressions."""
    def agg(_):
        return {
            "nuit_es": {n: {"date": "", "sr": "", "allees_str": str(n),
                            "eeg": 100, "rails_es": 0, "sa": 0, "cam": 0}
                        for n in range(1, 8)},
            "nuit_cam": {},
            "totals_by_nuit": {n: {"eeg": 100, "cam": 0, "sa": 0}
                                for n in range(1, 8)},
            "dates_map": {str(n): "" for n in range(1, 8)},
            "weeks": [1, 2, 1, 2, 1],  # 5 semaines
            "all_nights": list(range(1, 8)),
            "cam_nights": [],
            "hide_sa_mag": False,
        }
    d = _minimal_dataset()
    data = pptx_export.build_pptx(
        d, aggregate_fn=agg, recap_rows=[],
        summary={"allees": [], "totals": {}}, detail_cam_rows=[],
    )
    _check_no_zip_duplicates(data)
