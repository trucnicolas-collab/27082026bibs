#!/usr/bin/env python3
"""
Backend test suite for VT/Phasage Carrefour app - Plan wifi feature testing.
Tests ONLY the new "Plan wifi" backend endpoints.
"""
import requests
import io
import os
from pathlib import Path
from PIL import Image
from openpyxl import Workbook
from pptx import Presentation

# Backend URL from frontend/.env
BACKEND_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://go-lang-43.preview.emergentagent.com")
API_BASE = f"{BACKEND_URL}/api"

# Test credentials from /app/memory/test_credentials.md
ADMIN_EMAIL = "admin@vusion.local"
ADMIN_PASSWORD = "admin123"

# Global state
auth_token = None
upload_id = None


def login():
    """Login as admin and get JWT token."""
    global auth_token
    print("\n=== LOGIN ===")
    resp = requests.post(
        f"{API_BASE}/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Login failed: {resp.text}"
    data = resp.json()
    print(f"Logged in as: {data.get('email')}")
    
    # Extract token from cookies or response
    if "access_token" in resp.cookies:
        auth_token = resp.cookies["access_token"]
    else:
        # Try to get from response body if available
        auth_token = data.get("access_token", "")
    
    # For Bearer token auth, we'll use cookies
    return resp.cookies


def create_synthetic_excel():
    """Create a synthetic Excel file with required columns and EEG data."""
    print("\n=== CREATING SYNTHETIC EXCEL ===")
    wb = Workbook()
    ws = wb.active
    
    # Header row - EXACT column names required
    headers = ["Secteur", "Rayon", "N° allée", "Type", "Référence", "Désignation", "Quantité"]
    ws.append(headers)
    
    # Add ~8 rows with EEG data
    rows = [
        ["NAL", "Conserves", "1", "EEG", "15024", "ES 1.5 (noir)", 50],
        ["NAL", "Conserves", "2", "EEG", "17740", "ES 2.1 (noir)", 30],
        ["NAL", "Liquides", "3", "EEG", "15024", "ES 1.5 (noir)", 40],
        ["PGC", "Épicerie", "4", "EEG", "15910", "SA 2.1 (noir)", 60],
        ["PGC", "Épicerie", "5", "EEG", "15551", "SA 1.5 (noir)", 25],
        ["NAL", "Conserves", "1", "Rail", "16957", "1187mm noir", 10],
        ["NAL", "Conserves", "2", "Rail", "15507", "1240mm noir", 8],
        ["PGC", "Épicerie", "4", "Caméra", "11892", "Caméra (blanche)", 5],
    ]
    for row in rows:
        ws.append(row)
    
    # Save to bytes
    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    print(f"Created Excel with {len(rows)} data rows")
    return excel_bytes


def upload_excel(cookies):
    """Upload synthetic Excel and get upload_id."""
    global upload_id
    print("\n=== UPLOADING EXCEL ===")
    
    excel_bytes = create_synthetic_excel()
    files = {"file": ("test_inventory.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    
    resp = requests.post(
        f"{API_BASE}/upload-excel",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Upload failed: {resp.text}"
    
    data = resp.json()
    upload_id = data["upload_id"]
    print(f"Upload ID: {upload_id}")
    print(f"Filename: {data['filename']}")
    print(f"Row count: {data['row_count']}")
    return upload_id


def create_test_image(width=800, height=600, color="blue", format="PNG"):
    """Create a test image using PIL."""
    img = Image.new("RGB", (width, height), color=color)
    img_bytes = io.BytesIO()
    img.save(img_bytes, format=format)
    img_bytes.seek(0)
    return img_bytes


def test_wifi_plan_upload(cookies, upload_id):
    """Test 1: Upload wifi plans (PNG, JPG, max 2, error cases)."""
    print("\n=== TEST 1: WIFI PLAN UPLOAD ===")
    
    # Test 1a: Upload first PNG image
    print("\n1a. Upload first PNG image")
    img1 = create_test_image(800, 600, "blue", "PNG")
    files = {"file": ("plan1.png", img1, "image/png")}
    resp = requests.post(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Upload PNG failed: {resp.text}"
    data = resp.json()
    print(f"Response: {data}")
    assert data["count"] == 1, f"Expected count=1, got {data['count']}"
    assert data["max"] == 2, f"Expected max=2, got {data['max']}"
    assert len(data["plans"]) == 1, f"Expected 1 plan, got {len(data['plans'])}"
    plan1_id = data["plans"][0]["plan_id"]
    print(f"✓ First plan uploaded: {plan1_id}")
    
    # Test 1b: Upload second JPG image
    print("\n1b. Upload second JPG image")
    img2 = create_test_image(1024, 768, "red", "JPEG")
    files = {"file": ("plan2.jpg", img2, "image/jpeg")}
    resp = requests.post(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Upload JPG failed: {resp.text}"
    data = resp.json()
    print(f"Response: {data}")
    assert data["count"] == 2, f"Expected count=2, got {data['count']}"
    assert len(data["plans"]) == 2, f"Expected 2 plans, got {len(data['plans'])}"
    plan2_id = data["plans"][1]["plan_id"]
    print(f"✓ Second plan uploaded: {plan2_id}")
    
    # Test 1c: Try to upload third image (should fail - max 2)
    print("\n1c. Try to upload third image (should fail)")
    img3 = create_test_image(640, 480, "green", "PNG")
    files = {"file": ("plan3.png", img3, "image/png")}
    resp = requests.post(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 400, f"Expected 400 for max exceeded, got {resp.status_code}"
    print(f"Error message: {resp.json().get('detail')}")
    print("✓ Correctly rejected third image (max 2)")
    
    # Test 1d: Try to upload non-image file (should fail)
    print("\n1d. Try to upload non-image file (should fail)")
    txt_content = b"This is a text file, not an image"
    files = {"file": ("test.txt", io.BytesIO(txt_content), "text/plain")}
    resp = requests.post(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 400, f"Expected 400 for unsupported format, got {resp.status_code}"
    print(f"Error message: {resp.json().get('detail')}")
    print("✓ Correctly rejected non-image file")
    
    return plan1_id, plan2_id


def test_wifi_plan_list(cookies, upload_id):
    """Test 2: List wifi plans."""
    print("\n=== TEST 2: LIST WIFI PLANS ===")
    
    resp = requests.get(
        f"{API_BASE}/dataset/{upload_id}/wifi-plans",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"List failed: {resp.text}"
    
    data = resp.json()
    print(f"Response: {data}")
    assert data["count"] == 2, f"Expected 2 plans, got {data['count']}"
    assert len(data["plans"]) == 2, f"Expected 2 plans in list, got {len(data['plans'])}"
    
    # Check plan structure
    for i, plan in enumerate(data["plans"]):
        print(f"\nPlan {i}:")
        print(f"  plan_id: {plan['plan_id']}")
        print(f"  filename: {plan['filename']}")
        print(f"  content_type: {plan['content_type']}")
        print(f"  position: {plan['position']}")
        assert "plan_id" in plan
        assert "filename" in plan
        assert "content_type" in plan
        assert "position" in plan
        assert plan["position"] == i, f"Expected position {i}, got {plan['position']}"
    
    print("✓ List endpoint working correctly")
    return data["plans"]


def test_wifi_plan_get(cookies, upload_id, plan_id, expected_size_min=1000):
    """Test 3: Get single wifi plan image."""
    print(f"\n=== TEST 3: GET WIFI PLAN {plan_id} ===")
    
    resp = requests.get(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan/{plan_id}",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Get plan failed: {resp.text}"
    
    # Check Content-Type
    content_type = resp.headers.get("Content-Type", "")
    print(f"Content-Type: {content_type}")
    assert content_type in ["image/png", "image/jpeg"], f"Expected image content type, got {content_type}"
    
    # Check body is non-empty binary
    body = resp.content
    print(f"Body size: {len(body)} bytes")
    assert len(body) > expected_size_min, f"Expected body > {expected_size_min} bytes, got {len(body)}"
    
    # Verify it's a valid image
    try:
        img = Image.open(io.BytesIO(body))
        print(f"Image size: {img.size}")
        print(f"Image format: {img.format}")
        print("✓ Valid image retrieved")
    except Exception as e:
        raise AssertionError(f"Invalid image data: {e}")
    
    return body


def test_wifi_plan_delete(cookies, upload_id, plan_id):
    """Test 4: Delete wifi plan and verify re-indexing."""
    print(f"\n=== TEST 4: DELETE WIFI PLAN {plan_id} ===")
    
    resp = requests.delete(
        f"{API_BASE}/dataset/{upload_id}/wifi-plan/{plan_id}",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Delete failed: {resp.text}"
    
    data = resp.json()
    print(f"Response: {data}")
    assert data["count"] == 1, f"Expected count=1 after delete, got {data['count']}"
    assert len(data["plans"]) == 1, f"Expected 1 plan remaining, got {len(data['plans'])}"
    
    # Check remaining plan is re-indexed to position 0
    remaining_plan = data["plans"][0]
    print(f"Remaining plan: {remaining_plan}")
    assert remaining_plan["position"] == 0, f"Expected position 0, got {remaining_plan['position']}"
    print("✓ Plan deleted and positions re-indexed correctly")
    
    return remaining_plan["plan_id"]


def test_pptx_export_with_plans(cookies, upload_id, num_plans_expected):
    """Test 5: Export PPTX and verify wifi plan slides.
    
    CRITICAL VERIFICATION (per review request):
    - Count ALL slides whose ANY shape text contains "plan wifi" (case-insensitive)
    - For each such slide, count PICTURE shapes (shape.shape_type == 13)
    - 2 plans → EXACTLY 2 "Plan wifi" slides, EACH with EXACTLY 1 picture, CONSECUTIVE positions
    - 1 plan → EXACTLY 1 "Plan wifi" slide with EXACTLY 1 picture
    - 0 plans → EXACTLY 1 "Plan wifi" slide with 0 pictures
    - NO third "Plan wifi" slide anywhere (this was the bug)
    """
    print(f"\n=== TEST 5: EXPORT PPTX (expecting {num_plans_expected} wifi plan slides) ===")
    
    resp = requests.get(
        f"{API_BASE}/export-pptx/{upload_id}",
        cookies=cookies,
        timeout=60
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 200, f"Export failed: {resp.text}"
    
    # Check Content-Type
    content_type = resp.headers.get("Content-Type", "")
    print(f"Content-Type: {content_type}")
    assert "presentationml.presentation" in content_type, f"Expected PPTX content type, got {content_type}"
    
    # Check body is non-empty
    body = resp.content
    print(f"PPTX size: {len(body) / (1024*1024):.1f} MB")
    assert len(body) > 10000, f"PPTX too small: {len(body)} bytes"
    
    # Load PPTX with python-pptx and verify after reopen (this is where the bug manifested)
    try:
        prs = Presentation(io.BytesIO(body))
        print(f"Total slides in deck: {len(prs.slides)}")
        
        # Find ALL "Plan wifi" slides (ANY shape text contains "plan wifi", case-insensitive)
        wifi_slides_info = []
        for i, slide in enumerate(prs.slides):
            has_wifi_text = False
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.lower()
                    if "plan wifi" in text:
                        has_wifi_text = True
                        break
            
            if has_wifi_text:
                # Count PICTURE shapes (shape_type == 13)
                pictures = [s for s in slide.shapes if s.shape_type == 13]
                wifi_slides_info.append({
                    "position": i,
                    "num_pictures": len(pictures)
                })
                print(f"  Slide {i}: 'Plan wifi' text found, {len(pictures)} picture(s)")
        
        num_wifi_slides = len(wifi_slides_info)
        print(f"\n>>> OBSERVED: {num_wifi_slides} 'Plan wifi' slide(s) in deck")
        
        # PRIMARY VERIFICATION: Exact count based on num_plans_expected
        if num_plans_expected == 2:
            # 2 plans → EXACTLY 2 "Plan wifi" slides
            assert num_wifi_slides == 2, \
                f"FAIL: Expected EXACTLY 2 'Plan wifi' slides for 2 plans, found {num_wifi_slides}"
            print("✓ PASS: Exactly 2 'Plan wifi' slides (no phantom slide)")
            
            # Each slide must have EXACTLY 1 picture
            for info in wifi_slides_info:
                assert info["num_pictures"] == 1, \
                    f"FAIL: Slide {info['position']} has {info['num_pictures']} pictures, expected 1"
            print("✓ PASS: Each 'Plan wifi' slide has exactly 1 picture")
            
            # Slides must be CONSECUTIVE
            positions = [info["position"] for info in wifi_slides_info]
            assert positions[1] == positions[0] + 1, \
                f"FAIL: 'Plan wifi' slides not consecutive: positions {positions}"
            print(f"✓ PASS: 'Plan wifi' slides are consecutive (positions {positions})")
            
        elif num_plans_expected == 1:
            # 1 plan → EXACTLY 1 "Plan wifi" slide with EXACTLY 1 picture
            assert num_wifi_slides == 1, \
                f"FAIL: Expected EXACTLY 1 'Plan wifi' slide for 1 plan, found {num_wifi_slides}"
            print("✓ PASS: Exactly 1 'Plan wifi' slide")
            
            assert wifi_slides_info[0]["num_pictures"] == 1, \
                f"FAIL: 'Plan wifi' slide has {wifi_slides_info[0]['num_pictures']} pictures, expected 1"
            print("✓ PASS: 'Plan wifi' slide has exactly 1 picture")
            
        elif num_plans_expected == 0:
            # 0 plans → EXACTLY 1 "Plan wifi" slide with 0 pictures (empty slide, no crash)
            assert num_wifi_slides == 1, \
                f"FAIL: Expected EXACTLY 1 'Plan wifi' slide for 0 plans, found {num_wifi_slides}"
            print("✓ PASS: Exactly 1 'Plan wifi' slide (empty)")
            
            assert wifi_slides_info[0]["num_pictures"] == 0, \
                f"FAIL: 'Plan wifi' slide has {wifi_slides_info[0]['num_pictures']} pictures, expected 0"
            print("✓ PASS: 'Plan wifi' slide has 0 pictures (empty)")
        
        # Verify the "Commandes" slide (title "8") still exists
        print("\nVerifying rest of deck is intact:")
        commandes_found = False
        for i, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text == "8" or "commandes" in text.lower():
                        commandes_found = True
                        print(f"  ✓ Found 'Commandes' slide at position {i}")
                        break
            if commandes_found:
                break
        
        assert commandes_found, "FAIL: Could not find 'Commandes' slide - deck may be corrupted"
        
        # Verify total slide count is reasonable (~18 slides for a typical deck)
        assert 15 <= len(prs.slides) <= 25, \
            f"FAIL: Total slide count {len(prs.slides)} seems unreasonable (expected ~18)"
        print(f"  ✓ Total slide count reasonable: {len(prs.slides)}")
        
        print(f"\n✓✓✓ PPTX export with {num_plans_expected} wifi plan(s) VERIFIED SUCCESSFULLY ✓✓✓")
        
    except AssertionError:
        raise
    except Exception as e:
        raise AssertionError(f"Failed to parse PPTX: {e}")
    
    return body


def test_error_cases(cookies):
    """Test 6: Error cases with non-existent upload_id."""
    print("\n=== TEST 6: ERROR CASES (non-existent upload_id) ===")
    
    fake_id = "00000000-0000-0000-0000-000000000000"
    
    # Test POST with fake ID
    print("\n6a. POST wifi-plan with fake upload_id")
    img = create_test_image(640, 480, "yellow", "PNG")
    files = {"file": ("test.png", img, "image/png")}
    resp = requests.post(
        f"{API_BASE}/dataset/{fake_id}/wifi-plan",
        files=files,
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 404, f"Expected 404, got {resp.status_code}"
    print("✓ POST correctly returns 404 for non-existent ID")
    
    # Test GET list with fake ID
    print("\n6b. GET wifi-plans with fake upload_id")
    resp = requests.get(
        f"{API_BASE}/dataset/{fake_id}/wifi-plans",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 404, f"Expected 404, got {resp.status_code}"
    print("✓ GET list correctly returns 404 for non-existent ID")
    
    # Test GET single with fake ID
    print("\n6c. GET wifi-plan/{plan_id} with fake upload_id")
    resp = requests.get(
        f"{API_BASE}/dataset/{fake_id}/wifi-plan/fake-plan-id",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 404, f"Expected 404, got {resp.status_code}"
    print("✓ GET single correctly returns 404 for non-existent ID")
    
    # Test DELETE with fake ID
    print("\n6d. DELETE wifi-plan with fake upload_id")
    resp = requests.delete(
        f"{API_BASE}/dataset/{fake_id}/wifi-plan/fake-plan-id",
        cookies=cookies,
        timeout=30
    )
    print(f"Status: {resp.status_code}")
    assert resp.status_code == 404, f"Expected 404, got {resp.status_code}"
    print("✓ DELETE correctly returns 404 for non-existent ID")
    
    print("\n✓ All error cases handled correctly")


def main():
    """Main test execution."""
    print("=" * 70)
    print("WIFI PLAN BACKEND TEST SUITE")
    print("=" * 70)
    
    try:
        # Setup
        cookies = login()
        upload_id = upload_excel(cookies)
        
        # Test 1: Upload wifi plans (2 images, test max limit and format validation)
        plan1_id, plan2_id = test_wifi_plan_upload(cookies, upload_id)
        
        # Test 2: List wifi plans
        plans = test_wifi_plan_list(cookies, upload_id)
        
        # Test 3: Get individual wifi plans
        img1_data = test_wifi_plan_get(cookies, upload_id, plan1_id)
        img2_data = test_wifi_plan_get(cookies, upload_id, plan2_id)
        
        # Test 5a: Export PPTX with 2 plans
        pptx_2plans = test_pptx_export_with_plans(cookies, upload_id, num_plans_expected=2)
        
        # Test 4: Delete one plan
        remaining_plan_id = test_wifi_plan_delete(cookies, upload_id, plan1_id)
        
        # Test 5b: Export PPTX with 1 plan
        pptx_1plan = test_pptx_export_with_plans(cookies, upload_id, num_plans_expected=1)
        
        # Delete remaining plan
        print("\n=== DELETING REMAINING PLAN ===")
        resp = requests.delete(
            f"{API_BASE}/dataset/{upload_id}/wifi-plan/{remaining_plan_id}",
            cookies=cookies,
            timeout=30
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["count"] == 0, f"Expected count=0, got {data['count']}"
        print("✓ All plans deleted")
        
        # Test 5c: Export PPTX with 0 plans (should still work, 1 empty wifi slide)
        pptx_0plans = test_pptx_export_with_plans(cookies, upload_id, num_plans_expected=0)
        
        # Test 6: Error cases
        test_error_cases(cookies)
        
        print("\n" + "=" * 70)
        print("ALL TESTS PASSED ✓")
        print("=" * 70)
        print("\nSUMMARY:")
        print("✓ Upload wifi plans (PNG, JPG)")
        print("✓ Max 2 plans enforced")
        print("✓ Unsupported formats rejected")
        print("✓ List wifi plans")
        print("✓ Get individual wifi plan images")
        print("✓ Delete wifi plan with re-indexing")
        print("✓ PPTX export with 2 plans (2 slides with 1 picture each)")
        print("✓ PPTX export with 1 plan (1 slide with 1 picture)")
        print("✓ PPTX export with 0 plans (no crash)")
        print("✓ Error handling for non-existent IDs")
        
    except AssertionError as e:
        print(f"\n❌ TEST FAILED: {e}")
        return 1
    except Exception as e:
        print(f"\n❌ UNEXPECTED ERROR: {e}")
        import traceback
        traceback.print_exc()
        return 1
    
    return 0


if __name__ == "__main__":
    exit(main())
